"""Strict paragraph-based .docx template renderer for the starter MVP.

Supported tags (each control tag on its own paragraph):
  {{ full_name }}, {{ exp.company }}, etc.
  {%p for exp in experiences %} ... {%p endfor %}
Nested loops are supported. Runs must not split a placeholder.
"""
from __future__ import annotations
import copy
import re
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


FIELD = re.compile(r"{{\s*([a-zA-Z_]\w*(?:\.[a-zA-Z_]\w*)*)\s*}}")
START = re.compile(r"^\{%p\s+for\s+([a-zA-Z_]\w*)\s+in\s+([a-zA-Z_]\w*(?:\.[a-zA-Z_]\w*)*)\s*%}$")
END = re.compile(r"^\{%p\s+endfor\s*%}$")
TAG = re.compile(r"{{|{%")
MAX_ITERATIONS = 200

class TemplateError(ValueError):
    pass

def lookup(path: str, context: dict):
    value = context
    for part in path.split('.'):
        if not isinstance(value, dict) or part not in value:
            if path in ('exp.description', 'linkedin'):
                return ''  # Optional fields in older resume JSON payloads.
            raise TemplateError(f"Missing field: {path}")
        value = value[part]
    return value

def text_of(element):
    return ''.join(t.text or '' for t in element.iter(qn('w:t')))

BOLD_MARKUP = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)


def apply_bold_markup(node, value: str, regular_default: bool = False):
    """Turn **keywords** into actual bold Word runs, preserving template fonts.

    Handles placeholders surrounded by text in the same run, e.g.
    "• {{ bullet.text }}". Unmatched ** is left as ordinary text.
    """
    spans = []
    last = 0
    for match in BOLD_MARKUP.finditer(value):
        if match.start() > last:
            spans.append((value[last:match.start()], False))
        spans.append((match.group(1), True))
        last = match.end()
    if not spans:
        spans = [(value, False)]
    elif last < len(value):
        spans.append((value[last:], False))
    source_run = node.getparent()
    if source_run is None or source_run.tag != qn('w:r'):
        node.text = BOLD_MARKUP.sub(lambda m: m.group(1), value)
        return
    parent = source_run.getparent()
    if parent is None or len(list(source_run.iter(qn('w:t')))) != 1:
        node.text = BOLD_MARKUP.sub(lambda m: m.group(1), value)
        return
    position = parent.index(source_run)
    for text, is_bold in spans:
        # A skills_text category is separated by \n. Word requires a w:br
        # element for each break; embedding raw newlines in w:t is unreliable.
        for segment_index, segment in enumerate(text.split('\n')):
            if segment_index:
                br_run = copy.deepcopy(source_run)
                br_node = next(br_run.iter(qn('w:t')))
                br_run.replace(br_node, OxmlElement('w:br'))
                parent.insert(position, br_run)
                position += 1
            if not segment:
                continue
            new_run = copy.deepcopy(source_run)
            new_node = next(new_run.iter(qn('w:t')))
            new_node.text = segment
            if segment != segment.strip():
                new_node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            if is_bold or regular_default:
                # An uploaded Word template can specify bold on its placeholder
                # run or paragraph style. Explicitly neutralize it for summary
                # body text; only **selected keywords** should remain bold.
                rpr = set_child(new_run, 'w:rPr')
                b = set_child(rpr, 'w:b')
                b.set(qn('w:val'), '1' if is_bold else '0')
                bcs = set_child(rpr, 'w:bCs')
                bcs.set(qn('w:val'), '1' if is_bold else '0')
            parent.insert(position, new_run)
            position += 1
    parent.remove(source_run)


def replace_fields(element, context, category: str = ""):
    # Preserve the uploaded Word template's paragraph and run formatting.
    # Values can contain **important words**; they become real Word bold runs.
    for node in list(element.iter(qn('w:t'))):
        if not node.text:
            continue
        def convert(match):
            value = lookup(match.group(1), context)
            if isinstance(value, (dict, list)):
                raise TemplateError(f"Field is not scalar: {match.group(1)}")
            return '' if value is None else str(value)
        original = node.text
        rendered = FIELD.sub(convert, original)
        # Imported PDF text may contain layout-only line breaks inside a
        # single bullet/summary/description. Word justification stretches
        # manually broken lines edge-to-edge, sometimes spilling into margins.
        # Make those passages one ordinary wrapping paragraph instead.
        if category in ('summary', 'skills', 'experienceDescription', 'experienceBullets', 'education', 'body'):
            rendered = re.sub(r'\s*[\r\n]+\s*', ' ', rendered)
        if rendered != original and (category == 'summary' or '**' in rendered or '\n' in rendered):
            apply_bold_markup(node, rendered, regular_default=(category == 'summary'))
        else:
            node.text = rendered
    remaining = text_of(element)
    if TAG.search(remaining):
        raise TemplateError("Unsupported tag or placeholder split across Word formatting runs: " + remaining[:90])

# User-controlled typography is applied after field replacement to each cloned
# Word paragraph. Uploaded template margins/tables/bullets remain unchanged.
FONT_ALLOWLIST = {'Times New Roman', 'Arial', 'Calibri', 'Aptos', 'Georgia', 'Verdana'}
STYLE_KEYS = {
    'name', 'role', 'personal', 'summary', 'body',
    'skillsHeading', 'skills', 'experienceHeading', 'experience',
    'experienceDescription', 'experienceBullets',
    'educationHeading', 'education',
}
# Mapping both old single-style payloads and the new finer-grained style controls.
LEGACY_STYLE = {
    'skillsHeading': 'skills', 'experienceHeading': 'experience',
    'educationHeading': 'education',
    'experienceDescription': 'body', 'experienceBullets': 'body',
}
SECTION_HEADINGS = {
    'SUMMARY': 'summaryHeading', 'PROFESSIONAL SUMMARY': 'summaryHeading',
    'SKILLS': 'skillsHeading', 'TECHNICAL SKILLS': 'skillsHeading',
    'SKILLS & EXPERTISE': 'skillsHeading',
    'EXPERIENCE': 'experienceHeading', 'WORK EXPERIENCE': 'experienceHeading',
    'PROFESSIONAL EXPERIENCE': 'experienceHeading',
    'EDUCATION': 'educationHeading', 'EDUCATION & TRAINING': 'educationHeading',
}

def which_style(original: str, parent_context: dict) -> str:
    text = original.strip()
    if text.upper() in SECTION_HEADINGS:
        return SECTION_HEADINGS[text.upper()]
    if '{{' not in original:
        return 'body'
    if '{{ full_name' in original: return 'name'
    if '{{ job_title' in original: return 'role'
    if any('{{ '+x in original for x in ('email', 'phone', 'location', 'linkedin')): return 'personal'
    if '{{ summary' in original: return 'summary'
    if any('{{ '+x in original for x in ('skills_text', 'category.')): return 'skills'
    if '{{ exp.description' in original: return 'experienceDescription'
    if '{{ bullet.' in original: return 'experienceBullets'
    if '{{ exp.' in original: return 'experience'
    if '{{ ed.' in original: return 'education'
    return 'body'

def set_child(parent, name: str):
    tag = qn(name)
    child = parent.find(tag)
    if child is None:
        child = OxmlElement(name)
        # WordprocessingML requires paragraph and run properties to precede
        # paragraph/run content. Appending w:rPr after w:t produces invalid
        # run ordering; LibreOffice can then ignore the requested font in PDF.
        if name in ('w:pPr', 'w:rPr'):
            parent.insert(0, child)
        else:
            parent.append(child)
    elif name in ('w:pPr', 'w:rPr') and parent.index(child) != 0:
        # Normalize malformed runs created by older Folio versions too.
        parent.remove(child)
        parent.insert(0, child)
    return child

def apply_paragraph_style(paragraph, category: str, styles: dict):
    if paragraph.tag != qn('w:p') or not isinstance(styles, dict):
        return
    style = styles.get(category)
    # Never let body styles leak into section headings. In particular, the
    # default justified Skills body must not justify TECHNICAL SKILLS.
    if not isinstance(style, dict) and not category.endswith('Heading'):
        style = styles.get(LEGACY_STYLE.get(category, ''))
    if not isinstance(style, dict):
        return
    # A user-uploaded Word template is the formatting authority by default.
    # The frontend marks its built-in visual defaults as disabled. Previously
    # those defaults silently replaced centered/bold headings with 11pt left.
    if style.get('enabled') is False:
        return
    # The Summary uses keyword-level bold (**word**) instead of whole-
    # paragraph bold. Old browser sessions may have saved an accidental
    # bold=True summary override; never let it turn the entire summary bold.
    # Paragraph typography (font, size, alignment, color) still applies.
    if category == 'summary':
        style = {**style, 'bold': None}
        if isinstance(style.get('changed'), list):
            style['changed'] = [key for key in style['changed'] if key != 'bold']
    # For newly edited styles, apply only the properties the user changed;
    # leave other formatting (including center alignment and bold) intact.
    changed = style.get('changed')
    if isinstance(changed, list):
        # A saved style can be enabled while its changed list is empty (for
        # example after migration from a prior GuruCorp version). Applying its
        # default font/size/alignment would overwrite the Word template.
        allowed = {'font', 'size', 'color', 'align', 'bold', 'lineSpacing', 'spaceAfter', 'leftIndent'}
        selected = set(changed).intersection(allowed)
        if not selected:
            return
        style = {k: v for k, v in style.items() if k in selected}
    font = style.get('font')
    size = style.get('size')
    color = style.get('color')
    align = style.get('align')
    bold = style.get('bold')
    line_spacing = style.get('lineSpacing')
    space_after = style.get('spaceAfter')
    left_indent = style.get('leftIndent')
    if font not in FONT_ALLOWLIST: font = None
    if type(size) not in (int,float) or not 7 <= size <= 36: size = None
    if not isinstance(color,str) or not re.fullmatch(r'#[0-9a-fA-F]{6}', color): color = None
    if align not in ('left','right','center','justify'): align = None
    if type(line_spacing) not in (int, float) or not 1 <= line_spacing <= 2.5: line_spacing = None
    if type(space_after) not in (int, float) or not 0 <= space_after <= 36: space_after = None
    if type(left_indent) not in (int, float) or not 0 <= left_indent <= 72: left_indent = None
    if line_spacing is not None or space_after is not None:
        ppr = set_child(paragraph, 'w:pPr')
        spacing = ppr.get_or_add_spacing()
        if line_spacing is not None:
            spacing.set(qn('w:line'), str(round(line_spacing * 240)))
            spacing.set(qn('w:lineRule'), 'auto')
        if space_after is not None:
            spacing.set(qn('w:after'), str(round(space_after * 20)))
    # Do not change Word's indentation unless the user explicitly sets it.
    if left_indent is not None and category in ('skills', 'experienceBullets'):
        ppr = set_child(paragraph, 'w:pPr')
        indentation = ppr.get_or_add_ind()
        indentation.set(qn('w:left'), str(round(left_indent * 20)))
        if category == 'experienceBullets' and indentation.get(qn('w:hanging')) is None:
            indentation.set(qn('w:hanging'), str(round(min(12, left_indent) * 20)))
    if align:
        ppr = set_child(paragraph, 'w:pPr')
        jc = ppr.get_or_add_jc()
        jc.set(qn('w:val'), 'both' if align == 'justify' else align)
    for run in paragraph.iter(qn('w:r')):
        rpr = set_child(run, 'w:rPr')
        if font:
            fonts = set_child(rpr, 'w:rFonts')
            for attr in ('ascii','hAnsi','eastAsia','cs'):
                fonts.set(qn('w:'+attr), font)
        if size:
            sz = set_child(rpr, 'w:sz')
            sz.set(qn('w:val'), str(round(size*2)))
            szcs = set_child(rpr, 'w:szCs')
            szcs.set(qn('w:val'), str(round(size*2)))
        if color:
            c = set_child(rpr, 'w:color')
            c.set(qn('w:val'), color[1:])
        # Do not remove per-word bold supplied by **text** markup.
        # Preserve explicit rich-text **keyword** bold in summary/bullets.
        # For headings, an explicit false should override template bold.
        if bold is True or (bold is False and (category.endswith('Heading') or rpr.find(qn('w:b')) is None)):
            b = set_child(rpr, 'w:b')
            b.set(qn('w:val'), '1' if bold else '0')

def keep_bullets_inside_page(paragraph, category: str, styles: dict, source_text: str = "",
                             printable_width_twips: int | None = None):
    """Keep justified body paragraphs inside the printable Word page.

    Full justification is useful for multi-line body paragraphs, but not
    for title/date rows containing a right-aligned tab: justify + tab can
    push dates beyond the effective text boundary in LibreOffice. These
    rows use a right-aligned tab with left-aligned paragraph text instead.
    Normal body paragraphs remain justified and get a small safe right inset.
    """
    if paragraph.tag != qn('w:p'):
        return
    if category not in ('summary', 'skills', 'experience',
                        'experienceDescription', 'experienceBullets', 'education', 'body'):
        return
    ppr = set_child(paragraph, 'w:pPr')
    # Date rows must use one explicit right-aligned tab, not stretched spaces.
    is_date_row = ('{{ exp.dates' in source_text or '{{ ed.dates' in source_text)
    has_tab = any(node.tag == qn('w:tab') for node in paragraph.iter())
    if is_date_row and has_tab:
        jc = ppr.get_or_add_jc()
        jc.set(qn('w:val'), 'left')
        # Clamp right-aligned tab stops which some exported templates place
        # beyond the available width. Leave a 0.10in inset from the right edge.
        tabs = ppr.find(qn('w:tabs'))
        if tabs is not None and printable_width_twips:
            for tab in tabs:
                if tab.get(qn('w:val')) != 'right':
                    continue
                try:
                    pos = int(tab.get(qn('w:pos'), '0'))
                except ValueError:
                    continue
                safe_pos = max(720, printable_width_twips - 144)
                if pos > safe_pos:
                    tab.set(qn('w:pos'), str(safe_pos))
        return
    # Keep a physical buffer inside the right page margin for justified text.
    # Never erase a larger intentional inset set in the Word template.
    indentation = ppr.get_or_add_ind()
    existing = indentation.get(qn('w:right'))
    try:
        value = int(existing) if existing is not None else 0
    except ValueError:
        value = 0
    minimum = 216 if category == 'experienceBullets' else 144
    indentation.set(qn('w:right'), str(max(minimum, value)))
    # Do not force alignment of the last line. Word's 'both' justification
    # automatically leaves the final line left-aligned when no hard break exists.

def expand_paragraphs(elements, context):
    output = []
    index = 0
    while index < len(elements):
        current = elements[index]
        content = text_of(current).strip()
        start = START.fullmatch(content)
        if END.fullmatch(content):
            raise TemplateError('Unexpected endfor')
        if start:
            variable, sequence_path = start.groups()
            nesting, end = 1, index + 1
            while end < len(elements):
                tag = text_of(elements[end]).strip()
                if START.fullmatch(tag): nesting += 1
                elif END.fullmatch(tag):
                    nesting -= 1
                    if nesting == 0: break
                end += 1
            if nesting:
                raise TemplateError('Loop missing endfor')
            values = lookup(sequence_path, context)
            if not isinstance(values, list) or len(values) > MAX_ITERATIONS:
                raise TemplateError('Loops require a list of at most 200 items')
            for value in values:
                if not isinstance(value, dict):
                    raise TemplateError('Loop items must be objects')
                scope = {**context, variable: value}
                output.extend(expand_paragraphs(elements[index + 1:end], scope))
            index = end + 1
            continue
        # Omit an empty optional company description to avoid extra blank lines.
        if content == '{{ ed.location }}' and not str(context.get('ed', {}).get('location') or '').strip():
            index += 1
            continue
        if content == '{{ exp.description }}' and not context.get('exp', {}).get('description', '').strip():
            index += 1
            continue
        # Legacy one-field Skills templates store every category in a single
        # newline-separated string. Use one real paragraph per category so
        # justified layout does not stretch manually broken lines.
        if content == '{{ skills_text }}' and '\n' in str(context.get('skills_text') or ''):
            for skill_line in str(context['skills_text']).splitlines():
                if not skill_line.strip():
                    continue
                item_context = {**context, 'skills_text': skill_line.strip()}
                new_line = copy.deepcopy(current)
                replace_fields(new_line, item_context, 'skills')
                styles = context.get('_styles', {})
                apply_paragraph_style(new_line, 'skills', styles)
                keep_bullets_inside_page(new_line, 'skills', styles, content,
                                         context.get('_printable_width_twips'))
                output.append(new_line)
            index += 1
            continue
        new = copy.deepcopy(current)
        category = which_style(content, context)
        replace_fields(new, context, category)
        styles = context.get('_styles', {})
        apply_paragraph_style(new, category, styles)
        keep_bullets_inside_page(new, category, styles, content,
                                 context.get('_printable_width_twips'))
        output.append(new)
        index += 1
    return output

def add_missing_description_placeholders(elements, context):
    """Support existing uploaded templates created before exp.description existed.

    Only add a company description inside a recognized `experiences` block,
    immediately after the paragraph containing `{{ exp.company }}`. Do not
    duplicate a template that already defines the description placeholder.
    """
    experiences = context.get('experiences', [])
    if not isinstance(experiences, list) or not any(
        isinstance(exp, dict) and str(exp.get('description') or '').strip()
        for exp in experiences
    ):
        return elements

    # A deliberate description placeholder in any location takes priority.
    if any('exp.description' in FIELD.findall(text_of(el)) for el in elements):
        return elements

    updated = []
    stack = []
    for element in elements:
        text = text_of(element).strip()
        start = START.fullmatch(text)
        if start:
            stack.append(start.group(2))
        elif END.fullmatch(text):
            if stack:
                stack.pop()
        updated.append(element)
        if ('experiences' not in stack or element.tag != qn('w:p')
                or 'exp.company' not in FIELD.findall(text)):
            continue

        # Copy company typography while explicitly clearing its bold and
        # bullet numbering; descriptions are always ordinary paragraphs.
        new_p = OxmlElement('w:p')
        ppr = element.find(qn('w:pPr'))
        if ppr is not None:
            ppr = copy.deepcopy(ppr)
            numbering = ppr.find(qn('w:numPr'))
            if numbering is not None:
                ppr.remove(numbering)
            new_p.append(ppr)
        run = OxmlElement('w:r')
        existing_run = element.find(qn('w:r'))
        if existing_run is not None:
            rpr = existing_run.find(qn('w:rPr'))
            if rpr is not None:
                rpr = copy.deepcopy(rpr)
                for tag in ('w:b', 'w:bCs'):
                    old_bold = rpr.find(qn(tag))
                    if old_bold is not None:
                        rpr.remove(old_bold)
                run.append(rpr)
        t = OxmlElement('w:t')
        t.text = '{{ exp.description }}'
        run.append(t)
        new_p.append(run)
        updated.append(new_p)
    return updated


def add_missing_linkedin_placeholder(elements, context):
    """Add optional LinkedIn to legacy template contact lines, once.

    An explicit {{ linkedin }} placeholder always wins. This preserves
    uploaded templates while allowing existing browser-saved DOCX files
    to export the new LinkedIn field without re-uploading a template.
    """
    linkedin = context.get('linkedin')
    if not isinstance(linkedin, str) or not linkedin.strip():
        return elements
    # Template authors decide placement if LinkedIn is already present.
    if any('linkedin' in FIELD.findall(text_of(el)) for el in elements):
        return elements
    # Also avoid duplicating hard-coded LinkedIn addresses.
    if any('linkedin.com/' in text_of(el).lower() for el in elements):
        return elements

    target = None
    score = -1
    for index, element in enumerate(elements):
        if element.tag != qn('w:p'):
            continue
        text = text_of(element)
        fields = set(FIELD.findall(text))
        if not fields.intersection({'location', 'phone', 'email'}):
            continue
        # Select the most complete contact line, preferring email.
        rating = len(fields.intersection({'location', 'phone', 'email'})) + (3 if 'email' in fields else 0)
        if rating > score:
            target, score = index, rating
    if target is None:
        return elements  # No safe insertion point: leave layout unchanged.

    updated = list(elements)
    line = copy.deepcopy(updated[target])
    # Preserve contact font from the last text-containing run.
    run = None
    for candidate in line.iter(qn('w:r')):
        if any(t.text for t in candidate.iter(qn('w:t'))):
            run = candidate
    new_run = OxmlElement('w:r')
    if run is not None:
        properties = run.find(qn('w:rPr'))
        if properties is not None:
            new_run.append(copy.deepcopy(properties))
    text = OxmlElement('w:t')
    text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    separator = ' ' if text_of(line).rstrip().endswith('|') else ' | '
    text.text = separator + '{{ linkedin }}'
    new_run.append(text)
    line.append(new_run)
    updated[target] = line
    return updated


def render_docx(template: Path, output: Path, context: dict):
    doc = Document(str(template))
    body = doc.element.body
    children = list(body)
    # Preserve trailing sectPr. Keep all repeating directives in body paragraphs.
    section_properties = [c for c in children if c.tag == qn('w:sectPr')]
    content = [c for c in children if c.tag != qn('w:sectPr')]
    for el in content:
        if el.tag != qn('w:p') and ('{%p' in text_of(el)):
            raise TemplateError('Repeating sections inside tables are not supported in v1')
    content = add_missing_description_placeholders(content, context)
    content = add_missing_linkedin_placeholder(content, context)
    # Calculate actual printable width from the uploaded Word template.
    # Do not assume every template is US Letter with 0.5in margins.
    page_widths = [int(section.page_width.twips - section.left_margin.twips - section.right_margin.twips)
                   for section in doc.sections if section.page_width and section.left_margin and section.right_margin]
    render_context = {**context, '_printable_width_twips': min(page_widths) if page_widths else 9360}
    rendered = expand_paragraphs(content, render_context)
    # Prevent LibreOffice/Word from stretching lines created by Shift+Enter
    # in third-party templates (including exports from Google Docs).
    settings_el = doc.settings.element
    compat = settings_el.find(qn('w:compat'))
    if compat is None:
        compat = OxmlElement('w:compat')
        settings_el.append(compat)
    if compat.find(qn('w:doNotExpandShiftReturn')) is None:
        compat.append(OxmlElement('w:doNotExpandShiftReturn'))
    for item in children: body.remove(item)
    for item in rendered + section_properties: body.append(item)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output

# Field names understood by the current resume editor. The renderer itself is
# intentionally more generic so advanced templates may introduce extra fields
# when the input JSON is extended in a future version.
SUPPORTED_FIELDS = {
    'full_name', 'job_title', 'email', 'phone', 'location', 'linkedin',
    'summary', 'skills_text', 'category.name', 'category.skills_text',
    'exp.job_title', 'exp.company', 'exp.dates', 'exp.description', 'bullet.text',
    'ed.degree', 'ed.institution', 'ed.location', 'ed.dates',
}
SUPPORTED_SEQUENCES = {'experiences', 'exp.bullets', 'education', 'skill_categories'}

def inspect_docx(template: Path) -> dict:
    """Return detected template placeholders or a descriptive validation error."""
    doc = Document(str(template))
    for section in doc.sections:
        for part in (section.header, section.footer):
            if TAG.search(text_of(part._element)):
                raise TemplateError('Placeholders inside headers and footers are not supported in v1')
    elements = list(doc.element.body)
    fields = set()
    sequences = set()
    stack = []
    for element in elements:
        content = text_of(element).strip()
        start = START.fullmatch(content)
        end = END.fullmatch(content)
        if start:
            if element.tag != qn('w:p'):
                raise TemplateError('Repeating blocks must be in Word body paragraphs')
            var, path = start.groups()
            if path not in SUPPORTED_SEQUENCES:
                raise TemplateError(f'Unsupported repeating section: {path}')
            if (path == 'experiences' and var != 'exp') or (path == 'education' and var != 'ed') or (path == 'exp.bullets' and var != 'bullet') or (path == 'skill_categories' and var != 'category'):
                raise TemplateError(f'Unsupported loop variable for {path}')
            sequences.add(path)
            stack.append(path)
            continue
        if end:
            if not stack: raise TemplateError('Unexpected endfor')
            stack.pop()
            continue
        if '{%p' in content:
            raise TemplateError('Unsupported or misplaced repeating directive: ' + content[:70])
        if element.tag != qn('w:p') and '{%p' in content:
            raise TemplateError('Repeating blocks inside tables are not supported')
        for match in FIELD.finditer(content):
            path = match.group(1)
            if path not in SUPPORTED_FIELDS:
                raise TemplateError(f'Unsupported placeholder: {path}')
            fields.add(path)
        # Ensure tokens exist entirely within individual XML text nodes.
        for match in FIELD.finditer(content):
            if not any(match.group(0) in (node.text or '') for node in element.iter(qn('w:t'))):
                raise TemplateError('Placeholder split across Word formatting runs: ' + match.group(0))
        if TAG.search(FIELD.sub('', content)):
            raise TemplateError('Malformed or unsupported placeholder: ' + content[:70])
    if stack:
        raise TemplateError('Repeating block is missing an endfor')
    if not fields:
        raise TemplateError('No supported placeholders found in the Word template')
    return { 'fields': sorted(fields), 'loops': sorted(sequences), 'valid': True }
