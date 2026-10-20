import json
from pathlib import Path
from docx import Document
from app.engine import inspect_docx, render_docx
from app.importer import ResumeLine, parse_lines

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / 'examples' / 'Diego_Style_Category_Bullets_Template.docx'


def test_description_rendered_as_plain_paragraph(tmp_path):
    assert 'exp.description' in inspect_docx(TEMPLATE)['fields']
    data = json.loads((ROOT / 'examples' / 'sample_resume.json').read_text())
    generated = tmp_path / 'resume.docx'
    render_docx(TEMPLATE, generated, data)
    paragraphs = Document(generated).paragraphs
    idx = next(i for i,p in enumerate(paragraphs) if p.text == data['experiences'][0]['description'])
    assert data['experiences'][0]['company'] in paragraphs[idx-1].text
    assert paragraphs[idx+1].text == data['experiences'][0]['bullets'][0]['text']
    assert paragraphs[idx].style.name != 'List Paragraph'
    assert paragraphs[idx]._p.pPr is None or paragraphs[idx]._p.pPr.numPr is None


def test_description_optional_in_old_resume_json(tmp_path):
    data = json.loads((ROOT / 'examples' / 'sample_resume.json').read_text())
    for exp in data['experiences']:
        exp.pop('description')
    out = tmp_path / 'resume.docx'
    render_docx(TEMPLATE, out, data)
    assert all('{{' not in p.text for p in Document(out).paragraphs)


def test_import_description_is_not_a_bullet():
    lines = [ResumeLine('Jane Doe'), ResumeLine('Engineer'), ResumeLine('SUMMARY'),
             ResumeLine('Engineer building software.'), ResumeLine('WORK EXPERIENCE'),
             ResumeLine('Senior Engineer  Jan 2021 – Present'), ResumeLine('Acme Inc'),
             ResumeLine('Builds SaaS software for hospitals.'),
             ResumeLine('Improved API speed by 30%', True)]
    out = parse_lines(lines)['resume']['experiences'][0]
    assert out['description'] == 'Builds SaaS software for hospitals.'
    assert out['bullets'] == [{'text': 'Improved API speed by 30%'}]
