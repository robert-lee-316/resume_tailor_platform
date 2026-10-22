"""Default paragraph justification must not affect section headings or contact lines."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.engine import render_docx


def test_justified_body_keeps_headings_independent(tmp_path):
    template, output = tmp_path / 'template.docx', tmp_path / 'rendered.docx'
    doc = Document()
    for line in ('{{ full_name }}','{{ job_title }}','{{ location }} | {{ email }}', 'SUMMARY',
                 '{{ summary }}','TECHNICAL SKILLS','{%p for category in skill_categories %}',
                 '{{ category.name }}: {{ category.skills_text }}','{%p endfor %}',
                 'WORK EXPERIENCE','{%p for exp in experiences %}','{{ exp.job_title }}',
                 '{{ exp.description }}','{%p for bullet in exp.bullets %}',
                 '{{ bullet.text }}','{%p endfor %}','{%p endfor %}',
                 'EDUCATION','{%p for ed in education %}','{{ ed.institution }} — {{ ed.degree }}',
                 '{%p endfor %}'):
        p = doc.add_paragraph(line)
        if line in ('SUMMARY','TECHNICAL SKILLS','WORK EXPERIENCE','EDUCATION'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(template)
    align = lambda: {'align':'justify','enabled':True,'changed':['align']}
    styles={key:align() for key in ('summary','body','skills','experience','experienceDescription','experienceBullets','education')}
    payload={
        'full_name':'Alex Morgan','job_title':'Developer','location':'Austin, TX','email':'alex@example.com',
        'summary':'Created scalable APIs.','skill_categories':[{'name':'Languages','skills_text':'Python, SQL'}],
        'experiences':[{'job_title':'Engineer','description':'Built secure services.','bullets':[{'text':'Reduced costs by 25%.'}]}],
        'education':[{'institution':'Sample University','degree':'B.S.'}], '_styles':styles,
    }
    render_docx(template,output,payload)
    paragraphs=Document(output).paragraphs
    def find(text):return next(p for p in paragraphs if p.text == text)
    for heading in ('SUMMARY','TECHNICAL SKILLS','WORK EXPERIENCE','EDUCATION'):
        assert find(heading).alignment == WD_ALIGN_PARAGRAPH.CENTER
    for body in ('Created scalable APIs.','Languages: Python, SQL','Engineer','Built secure services.',
                 'Reduced costs by 25%.','Sample University — B.S.'):
        assert find(body).alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, body
    for personal in ('Alex Morgan','Developer','Austin, TX | alex@example.com'):
        assert find(personal).alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
