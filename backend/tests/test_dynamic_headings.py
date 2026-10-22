"""Editable heading properties must reach DOCX without fixed section presets."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from app.engine import render_docx


def test_three_headings_accept_independent_dynamic_settings(tmp_path):
    source, target = tmp_path / 'template.docx', tmp_path / 'result.docx'
    doc = Document()
    doc.add_paragraph('{{ full_name }}')
    for title in ('SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
        paragraph = doc.add_paragraph(title)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.runs[0]
        run.bold = True
        run.font.size = Pt(15)
    doc.save(source)

    settings = {
        'skillsHeading': {'enabled': True, 'changed': ['font', 'size', 'bold', 'align', 'color'],
                          'font': 'Georgia', 'size': 19, 'bold': False, 'align': 'right', 'color': '#884422'},
        'experienceHeading': {'enabled': True, 'changed': ['size', 'align'],
                              'size': 12.5, 'align': 'left', 'bold': False},
        'educationHeading': {'enabled': False, 'changed': ['font', 'size', 'bold', 'align'],
                             'font': 'Verdana', 'size': 30, 'bold': False, 'align': 'right'},
    }
    render_docx(source, target, {'full_name': 'Example', '_styles': settings})
    paragraphs = {p.text: p for p in Document(target).paragraphs}
    skills = paragraphs['SKILLS']
    assert skills.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert skills.runs[0].font.name == 'Georgia'
    assert skills.runs[0].font.size.pt == 19
    assert skills.runs[0].bold is False
    assert skills.runs[0].font.color.rgb.__str__() == '884422'
    experience = paragraphs['WORK EXPERIENCE']
    assert experience.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert experience.runs[0].font.size.pt == 12.5
    assert experience.runs[0].bold is True  # still inherited
    education = paragraphs['EDUCATION']
    assert education.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert education.runs[0].font.size.pt == 15
    assert education.runs[0].bold is True
