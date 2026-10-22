"""Default justification must be enforced in generated DOCX, not only HTML."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from app.engine import render_docx


def make_template(path):
    doc = Document()
    for txt in (
        '{{ full_name }}', '{{ job_title }}', '{{ email }}',
        'SUMMARY', '{{ summary }}', 'TECHNICAL SKILLS',
        '{%p for category in skill_categories %}',
        '{{ category.name }}: {{ category.skills_text }}',
        '{%p endfor %}',
        'WORK EXPERIENCE', '{%p for exp in experiences %}',
        '{{ exp.job_title }}\t{{ exp.dates }}', '{{ exp.company }}',
        '{{ exp.description }}', '{%p for bullet in exp.bullets %}',
        '{{ bullet.text }}', '{%p endfor %}', '{%p endfor %}',
        'EDUCATION', '{%p for ed in education %}',
        '{{ ed.institution }} — {{ ed.degree }}', '{{ ed.location }}',
        '{%p endfor %}',
    ):
        paragraph = doc.add_paragraph(txt)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if txt in ('SUMMARY', 'TECHNICAL SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if '\t' in txt:
            paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6))
    doc.save(path)


def payload():
    return {
        'full_name': 'Diego Soto', 'job_title': 'Software Engineer',
        'email': 'diego@example.org',
        'summary': 'Developed web applications and cloud infrastructure using Python, React and AWS.',
        'skill_categories': [{'name': 'Frontend', 'skills_text': 'React, Next.js, TypeScript'}],
        'experiences': [{
            'job_title': 'Developer', 'company': 'Acme', 'dates': '2022 – Present',
            'description': 'Designs platform features across many customers.',
            'bullets': [{'text': 'Built secure applications using Node.js, AWS and PostgreSQL.'}],
        }],
        'education': [{'institution': 'University', 'degree': 'B.S.', 'location': 'New York'}],
    }


def test_exact_word_export_justifies_contents_without_frontend_style_payload(tmp_path):
    source, output = tmp_path/'template.docx', tmp_path/'out.docx'
    make_template(source)
    render_docx(source, output, payload())
    paragraphs = {p.text: p for p in Document(output).paragraphs if p.text.strip()}
    for text in (
        payload()['summary'], 'Frontend: React, Next.js, TypeScript',
        'Acme', payload()['experiences'][0]['description'],
        payload()['experiences'][0]['bullets'][0]['text'],
        'University — B.S.', 'New York',
    ):
        assert paragraphs[text].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, text
    for text in ('SUMMARY', 'TECHNICAL SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
        assert paragraphs[text].alignment == WD_ALIGN_PARAGRAPH.CENTER
    for text in ('Diego Soto', 'Software Engineer', 'diego@example.org'):
        assert paragraphs[text].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraphs['Developer\t2022 – Present'].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_explicit_user_alignment_is_not_overwritten(tmp_path):
    source, output = tmp_path/'template.docx', tmp_path/'out.docx'
    make_template(source)
    data=payload()
    data['_styles']={
        'summary': {'enabled': True, 'changed':['align'], 'align':'left'},
        'skills': {'enabled': True, 'changed':['size'], 'size': 12, 'align':'left'},
        'experienceDescription': {'enabled': False, 'changed':[], 'align':'justify'},
    }
    render_docx(source, output, data)
    pars={p.text:p for p in Document(output).paragraphs if p.text.strip()}
    assert pars[data['summary']].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert pars['Frontend: React, Next.js, TypeScript'].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert pars[data['experiences'][0]['description']].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert pars[data['experiences'][0]['bullets'][0]['text']].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
