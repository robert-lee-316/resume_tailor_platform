from io import BytesIO
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.engine import inspect_docx, render_docx
from app.template_generator import generate_template_from_resume
from app.importer import import_resume

client = TestClient(app)


def sample_resume(tmp_path):
    doc = Document()
    for text in (
        'Jane Smith', 'Senior Software Engineer', 'Denver, CO | jane@example.com | +1 555 120 0034',
        'SUMMARY', 'Experienced software engineer building APIs.',
        'TECHNICAL SKILLS', 'Languages: Python, Java', 'Frontend: React, Angular',
        'WORK EXPERIENCE', 'Senior Software Engineer   Jun 2022 – Present',
        'Acme Corp | Denver, CO', 'Building enterprise software solutions.',
    ):
        doc.add_paragraph(text)
    doc.add_paragraph('Improved response times 35%.', style='List Bullet')
    for text in ('EDUCATION', 'State University – Bachelor of Science', 'Denver, CO   Aug 2014 – May 2018'):
        doc.add_paragraph(text)
    path = tmp_path / 'uploaded.docx'
    doc.save(path)
    return path


def test_generates_valid_reusable_template_and_renders(tmp_path):
    source = sample_resume(tmp_path)
    dest = tmp_path / 'generated.docx'
    assert generate_template_from_resume(source, dest) == 'generated'
    detected = inspect_docx(dest)
    assert {'skill_categories', 'experiences', 'exp.bullets', 'education'} <= set(detected['loops'])
    assert {'exp.description', 'ed.location', 'linkedin'} <= set(detected['fields'])
    payload = import_resume(source)['resume']
    payload['skill_categories'] = [{'name': 'Languages', 'skills_text': 'Python, Java'}]
    generated = tmp_path / 'populated.docx'
    render_docx(dest, generated, payload)
    text = '\n'.join(p.text for p in Document(generated).paragraphs)
    assert 'Jane Smith' in text
    assert 'Improved response times 35%' in text
    assert '{{' not in text and '{%p' not in text


def test_valid_uploaded_template_is_kept_as_is(tmp_path):
    source = sample_resume(tmp_path)
    template = tmp_path / 'created.docx'
    generate_template_from_resume(source, template)
    target = tmp_path / 'copy.docx'
    assert generate_template_from_resume(template, target) == 'existing'
    assert template.read_bytes() == target.read_bytes()


def test_api_generates_downloadable_docx(tmp_path):
    source = sample_resume(tmp_path)
    with source.open('rb') as stream:
        result = client.post('/api/generate-template-from-resume', files={'file': ('jane.docx', stream,
          'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert result.status_code == 200, result.text
    assert result.headers['X-GuruCorp-Template-Mode'] == 'generated'
    output = tmp_path / 'downloaded.docx'
    output.write_bytes(result.content)
    assert inspect_docx(output)['valid']


def test_rejects_invalid_placeholder_file(tmp_path):
    document = Document()
    document.add_paragraph('{{ misspelled }}')
    bad = tmp_path / 'bad.docx'
    document.save(bad)
    with bad.open('rb') as stream:
        response = client.post('/api/generate-template-from-resume', files={'file': ('bad.docx', stream,
          'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 422
    assert 'invalid' in response.json()['detail'].lower()


def test_repairs_split_placeholder_runs_without_losing_document(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph('Contact: ')
    paragraph.add_run('{{ link')
    paragraph.add_run('edin }}')
    path = tmp_path / 'split.docx'
    doc.save(path)
    fixed = tmp_path / 'fixed.docx'
    assert generate_template_from_resume(path, fixed) == 'repaired'
    assert inspect_docx(fixed)['valid']
    assert 'Contact:' in Document(fixed).paragraphs[0].text
    assert '{{ linkedin }}' in Document(fixed).paragraphs[0].text
