from pathlib import Path
from app.engine import inspect_docx, TemplateError
from docx import Document
import pytest

ROOT = Path(__file__).resolve().parents[2]

def test_sample_template_inspection():
    result = inspect_docx(ROOT / 'examples' / 'sample_template.docx')
    assert result['valid'] is True
    assert 'full_name' in result['fields']
    assert 'exp.company' in result['fields']
    assert 'experiences' in result['loops']


def test_reject_unsupported_placeholder(tmp_path):
    doc = Document()
    doc.add_paragraph('{{ mystery }}')
    path = tmp_path / 'bad.docx'
    doc.save(path)
    with pytest.raises(TemplateError, match='Unsupported placeholder'):
        inspect_docx(path)


def test_reject_empty_template(tmp_path):
    doc = Document()
    doc.add_paragraph('Only a heading')
    path = tmp_path / 'bad.docx'
    doc.save(path)
    with pytest.raises(TemplateError, match='No supported placeholders'):
        inspect_docx(path)
