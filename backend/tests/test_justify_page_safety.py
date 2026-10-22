"""Regression checks for GuruCorp justified text and tabbed date rows."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from app.engine import render_docx


def test_justification_wraps_without_forced_breaks_and_stays_inset(tmp_path):
    template, output = tmp_path / 'source.docx', tmp_path / 'result.docx'
    doc = Document()
    doc.add_paragraph('SUMMARY')
    doc.add_paragraph('{{ summary }}')
    doc.add_paragraph('WORK EXPERIENCE')
    doc.add_paragraph('{%p for exp in experiences %}')
    p = doc.add_paragraph('{{ exp.job_title }}\t{{ exp.dates }}')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.8))
    doc.add_paragraph('{%p for bullet in exp.bullets %}')
    doc.add_paragraph('{{ bullet.text }}', style='List Bullet')
    doc.add_paragraph('{%p endfor %}')
    doc.add_paragraph('{%p endfor %}')
    doc.save(template)
    body_justify = {'enabled': True, 'changed':['align'], 'align':'justify'}
    render_docx(template, output, {
        'summary':'Experience with APIs and\\n production architecture.',
        'experiences':[{'job_title':'Software Engineer', 'dates':'2023 - Present',
             'bullets':[{'text':'Architected systems with reliable databases\\n and cloud storage for a large company.'}]}],
        '_styles':{'summary':body_justify,'experience':body_justify,'experienceBullets':body_justify}
    })
    paras=Document(output).paragraphs
    summary=next(p for p in paras if p.text.startswith('Experience with'))
    bullet=next(p for p in paras if p.text.startswith('Architected systems'))
    dates=next(p for p in paras if 'Software Engineer' in p.text)
    assert summary.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY
    assert bullet.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY
    assert dates.alignment==WD_ALIGN_PARAGRAPH.LEFT
    assert summary.paragraph_format.right_indent.pt >= 7
    assert bullet.paragraph_format.right_indent.pt >= 10
    assert not list(summary._p.iter(qn('w:br')))
    assert not list(bullet._p.iter(qn('w:br')))


def test_skills_legacy_newlines_become_separate_paragraphs(tmp_path):
    template, output = tmp_path / 'source.docx', tmp_path / 'result.docx'
    doc=Document()
    doc.add_paragraph('{{ skills_text }}')
    doc.save(template)
    render_docx(template, output, {'skills_text':'Frontend: React, Angular\nBackend: Python, Node.js',
         '_styles':{'skills':{'enabled':True,'changed':['align'],'align':'justify'}}})
    lines=[p for p in Document(output).paragraphs if p.text.strip()]
    assert [p.text for p in lines]==['Frontend: React, Angular', 'Backend: Python, Node.js']
    assert all(p.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY for p in lines)
