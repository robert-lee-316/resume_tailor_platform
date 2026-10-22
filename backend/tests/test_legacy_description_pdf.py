"""Regression: Exact PDF must show the company description with old templates."""
import io
import json
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient
from app.main import app
from app.engine import render_docx

ROOT = Path(__file__).resolve().parents[2]
DATA = json.loads((ROOT / 'examples' / 'sample_resume.json').read_text())
TEMPLATE = ROOT / 'examples' / 'Diego_Style_Category_Bullets_Template.docx'


def old_template(path):
    doc = Document(TEMPLATE)
    for paragraph in list(doc.paragraphs):
        if '{{ exp.description }}' in paragraph.text:
            paragraph._element.getparent().remove(paragraph._element)
    doc.save(path)
    return path


def test_legacy_docx_populates_company_description_not_bullet(tmp_path):
    legacy = old_template(tmp_path / 'legacy.docx')
    output = tmp_path / 'rendered.docx'
    render_docx(legacy, output, DATA)
    paragraphs = Document(output).paragraphs
    description = DATA['experiences'][0]['description']
    matches = [i for i, p in enumerate(paragraphs) if p.text == description]
    assert len(matches) == 1
    idx = matches[0]
    assert DATA['experiences'][0]['company'] in paragraphs[idx-1].text
    assert paragraphs[idx]._p.pPr is None or paragraphs[idx]._p.pPr.numPr is None


def test_current_docx_does_not_duplicate_company_description(tmp_path):
    output = tmp_path / 'rendered.docx'
    render_docx(TEMPLATE, output, DATA)
    fulltext = [p.text for p in Document(output).paragraphs]
    assert fulltext.count(DATA['experiences'][0]['description']) == 1


def test_pdf_api_with_old_template_shows_description(tmp_path):
    from pypdf import PdfReader
    legacy = old_template(tmp_path / 'legacy.docx')
    client = TestClient(app)
    with legacy.open('rb') as f:
        response = client.post('/api/render', files={'template': ('legacy.docx', f)},
            data={'resume_json': json.dumps(DATA), 'format': 'pdf'})
    assert response.status_code == 200, response.text
    extracted = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert DATA['experiences'][0]['description'] in extracted


def test_no_placeholder_added_for_empty_description(tmp_path):
    legacy = old_template(tmp_path / 'legacy.docx')
    data = {**DATA, 'experiences': [{**exp, 'description': ''} for exp in DATA['experiences']]}
    output = tmp_path / 'empty.docx'
    render_docx(legacy, output, data)
    assert all('exp.description' not in p.text for p in Document(output).paragraphs)
