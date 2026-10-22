"""Ensure legacy (already uploaded) Word templates include LinkedIn in Exact PDF."""
import io
import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from app.engine import inspect_docx, render_docx
from app.main import app

ROOT = Path(__file__).resolve().parents[2]
EXAMPLE = ROOT / 'examples' / 'Diego_Style_Category_Bullets_Template.docx'
DATA = json.loads((ROOT / 'examples' / 'sample_resume.json').read_text())


def legacy_template(path):
    doc = Document(EXAMPLE)
    for para in doc.paragraphs:
        for run in para.runs:
            run.text = run.text.replace(' | {{ linkedin }}', '').replace('{{ linkedin }}', '')
    doc.save(path)
    assert 'linkedin' not in inspect_docx(path)['fields']
    return path


def test_old_template_has_linkedin_in_docx(tmp_path):
    old = legacy_template(tmp_path / 'legacy.docx')
    output = tmp_path / 'filled.docx'
    render_docx(old, output, DATA)
    paras = [p.text for p in Document(output).paragraphs]
    assert sum('linkedin.com/in/' in s for s in paras) == 1
    contact = next(s for s in paras if DATA['email'] in s)
    assert DATA['linkedin'] in contact
    assert DATA['phone'] in contact


def test_template_with_placeholder_does_not_duplicate(tmp_path):
    output = tmp_path / 'filled.docx'
    render_docx(EXAMPLE, output, DATA)
    combined = '\n'.join(p.text for p in Document(output).paragraphs)
    assert combined.count(DATA['linkedin']) == 1


def test_empty_linkedin_does_not_add_pipe(tmp_path):
    old = legacy_template(tmp_path / 'legacy.docx')
    doc = Document(old)
    expected_contact = next(p.text for p in doc.paragraphs if '{{ email }}' in p.text)
    output = tmp_path / 'filled.docx'
    render_docx(old, output, {**DATA, 'linkedin': ''})
    contact = next(p.text for p in Document(output).paragraphs if DATA['email'] in p.text)
    assert contact.count('|') == expected_contact.count('|')


def test_legacy_exact_pdf_has_linkedin(tmp_path):
    from pypdf import PdfReader
    old = legacy_template(tmp_path / 'legacy.docx')
    with old.open('rb') as source:
        response = TestClient(app).post(
            '/api/render', files={'template': ('legacy.docx', source)},
            data={'resume_json': json.dumps(DATA), 'format': 'pdf'},
        )
    assert response.status_code == 200, response.text
    pdf_text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert DATA['linkedin'] in pdf_text
