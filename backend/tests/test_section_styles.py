"""Distinct sections, details and bullet formatting in generated Word exports."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from app.engine import render_docx


def make_template(path):
    doc = Document()
    for line in (
        '{{ full_name }}',
        'SKILLS',
        '{%p for category in skill_categories %}',
        '{{ category.name }}: {{ category.skills_text }}',
        '{%p endfor %}',
        'WORK EXPERIENCE',
        '{%p for exp in experiences %}',
        '{{ exp.job_title }}',
        '{{ exp.company }}',
        '{{ exp.description }}',
        '{%p for bullet in exp.bullets %}',
        '{{ bullet.text }}',
        '{%p endfor %}',
        '{%p endfor %}',
        'EDUCATION',
        '{%p for ed in education %}',
        '{{ ed.institution }} — {{ ed.degree }}',
        '{{ ed.location }}',
        '{{ ed.dates }}',
        '{%p endfor %}',
    ):
        doc.add_paragraph(line)
    doc.save(path)


def context(styles):
    return {
        'full_name': 'Alex Morgan',
        'skill_categories': [{'name': 'Frontend', 'skills_text': 'React, Next.js'}],
        'experiences': [{'job_title': 'Developer', 'company': 'Acme', 'description': 'Built products.',
                         'bullets': [{'text': 'Cut latency by 35%'}]}],
        'education': [{'institution': 'State University', 'degree': 'CS', 'location': 'Austin, TX', 'dates': '2014–2018'}],
        '_styles': styles,
    }


def test_independent_section_heading_and_content_styles(tmp_path):
    template, result = tmp_path / 'template.docx', tmp_path / 'output.docx'
    make_template(template)
    styles = {
        'skillsHeading': {'font': 'Georgia', 'size': 16, 'color': '#113355', 'align': 'center', 'bold': True},
        'skills': {'font': 'Arial', 'size': 9, 'color': '#224466', 'align': 'justify', 'lineSpacing': 1.2},
        'experienceHeading': {'font': 'Verdana', 'size': 14},
        'experience': {'font': 'Calibri', 'size': 11},
        'experienceDescription': {'font': 'Aptos', 'size': 9, 'spaceAfter': 7},
        'experienceBullets': {'font': 'Arial', 'size': 10, 'leftIndent': 30, 'lineSpacing': 1.4},
        'educationHeading': {'font': 'Georgia', 'size': 15, 'align': 'right'},
        'education': {'font': 'Verdana', 'size': 8},
    }
    render_docx(template, result, context(styles))
    rendered = Document(result)
    items = {p.text: p for p in rendered.paragraphs if p.text}
    assert items['SKILLS'].runs[0].font.name == 'Georgia'
    assert items['SKILLS'].runs[0].font.size.pt == 16
    assert items['SKILLS'].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert items['Frontend: React, Next.js'].runs[0].font.name == 'Arial'
    assert items['Frontend: React, Next.js'].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert items['WORK EXPERIENCE'].runs[0].font.name == 'Verdana'
    assert items['Developer'].runs[0].font.name == 'Calibri'
    assert items['Built products.'].runs[0].font.name == 'Aptos'
    spacing = items['Built products.']._p.pPr.find(qn('w:spacing'))
    assert spacing.get(qn('w:after')) == '140'
    assert items['Cut latency by 35%'].runs[0].font.name == 'Arial'
    assert items['Cut latency by 35%']._p.pPr.find(qn('w:ind')).get(qn('w:left')) == '600'
    assert items['Cut latency by 35%']._p.pPr.find(qn('w:spacing')).get(qn('w:line')) == '336'
    assert items['EDUCATION'].runs[0].font.name == 'Georgia'
    assert items['EDUCATION'].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert items['State University — CS'].runs[0].font.name == 'Verdana'
    assert items['Austin, TX'].runs[0].font.size.pt == 8


def test_old_style_keys_still_work(tmp_path):
    template, result = tmp_path / 'template.docx', tmp_path / 'result.docx'
    make_template(template)
    render_docx(template, result, context({
        'skills': {'font': 'Arial', 'size': 13},
        'experience': {'font': 'Georgia', 'size': 12},
        'education': {'font': 'Verdana', 'size': 11},
        'body': {'font': 'Calibri', 'size': 10},
    }))
    items = {p.text: p for p in Document(result).paragraphs if p.text}
    assert items['SKILLS'].runs[0].font.name == 'Arial'
    assert items['WORK EXPERIENCE'].runs[0].font.name == 'Georgia'
    assert items['EDUCATION'].runs[0].font.name == 'Verdana'
    assert items['Built products.'].runs[0].font.name == 'Calibri'
    assert items['Cut latency by 35%'].runs[0].font.name == 'Calibri'
