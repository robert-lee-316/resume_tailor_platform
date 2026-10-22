"""Compatibility tests for categorized skills in both old and new Word templates."""
from docx import Document
from docx.oxml.ns import qn
from app.engine import render_docx, inspect_docx


def test_legacy_skills_placeholder_preserves_category_line_breaks(tmp_path):
    template = tmp_path / "legacy.docx"
    output = tmp_path / "rendered.docx"
    doc = Document()
    doc.add_paragraph("SKILLS & EXPERTISE")
    doc.add_paragraph("{{ skills_text }}")
    doc.save(template)
    assert inspect_docx(template)["fields"] == ["skills_text"]
    render_docx(template, output, {"skills_text": "Languages: Python, Java\nCloud: AWS, Azure"})
    para = Document(output).paragraphs[1]
    assert para.text == "Languages: Python, Java\nCloud: AWS, Azure"
    assert len(list(para._p.iter(qn('w:br')))) == 1


def test_new_category_loop_repeats_each_category_in_own_paragraph(tmp_path):
    template = tmp_path / "loop.docx"
    output = tmp_path / "rendered.docx"
    doc = Document()
    for line in ("SKILLS", "{%p for category in skill_categories %}",
                 "{{ category.name }}: {{ category.skills_text }}", "{%p endfor %}"):
        doc.add_paragraph(line)
    doc.save(template)
    metadata = inspect_docx(template)
    assert metadata['valid']
    assert 'skill_categories' in metadata['loops']
    render_docx(template, output, {"skill_categories": [
        {"name": "Languages", "skills_text": "Python • Java"},
        {"name": "Backend", "skills_text": "FastAPI • Django"},
    ]})
    assert [p.text for p in Document(output).paragraphs] == [
        'SKILLS', 'Languages: Python • Java', 'Backend: FastAPI • Django'
    ]


def test_bullet_for_each_category_in_existing_word_template(tmp_path):
    template = tmp_path / 'skill-template.docx'
    output = tmp_path / 'skill-result.docx'
    doc = Document()
    doc.add_paragraph('SKILLS & EXPERTISE')
    doc.add_paragraph('{{ skills_text }}')
    doc.save(template)
    value = '• Languages: Python, Java\n• Cloud: AWS, Azure\n• Tools: Git'
    render_docx(template, output, {'skills_text': value})
    result = Document(output)
    assert result.paragraphs[1].text == value
    assert len(list(result.paragraphs[1]._p.iter(qn('w:br')))) == 2


def test_delivered_sample_template_uses_one_bullet_per_category(tmp_path):
    from pathlib import Path
    template = Path(__file__).resolve().parents[2] / 'examples' / 'sample_template.docx'
    assert 'skill_categories' in inspect_docx(template)['loops']
    dest = tmp_path / 'result.docx'
    render_docx(template, dest, {
        'full_name': 'Alex Morgan', 'job_title': 'Engineer', 'email': 'alex@example.com',
        'phone': '123', 'location': 'Austin', 'summary': 'Experienced engineer.',
        'experiences': [], 'education': [],
        'skill_categories': [
            {'name': 'Languages', 'skills_text': 'Python, JavaScript'},
            {'name': 'Backend', 'skills_text': 'FastAPI, Django'},
        ],
    })
    doc = Document(dest)
    bullets = [p for p in doc.paragraphs if p.text.startswith(('Languages:', 'Backend:'))]
    assert len(bullets) == 2
    assert all(p.style.name == 'List Bullet' for p in bullets)
    assert all(p.runs[0].bold for p in bullets)
    assert [p.text for p in bullets] == [
        'Languages: Python, JavaScript', 'Backend: FastAPI, Django'
    ]

def test_delivered_diego_template_categories_have_individual_list_numbers(tmp_path):
    from pathlib import Path
    template = Path(__file__).resolve().parents[2] / 'examples' / 'Diego_Style_Category_Bullets_Template.docx'
    assert 'skill_categories' in inspect_docx(template)['loops']
    dest = tmp_path / 'result.docx'
    render_docx(template, dest, {
        'full_name': 'Alex Morgan', 'job_title': 'Engineer', 'email': 'alex@example.com',
        'phone': '123', 'location': 'Austin', 'summary': 'Experienced engineer.',
        'experiences': [], 'education': [],
        'skill_categories': [
            {'name': 'Languages', 'skills_text': 'Python, JavaScript'},
            {'name': 'Backend', 'skills_text': 'FastAPI, Django'},
        ],
    })
    doc = Document(dest)
    bullets = [p for p in doc.paragraphs if p.text.startswith(('Languages:', 'Backend:'))]
    assert len(bullets) == 2
    assert all(p._p.xpath('./w:pPr/w:numPr') for p in bullets)
    assert all(p.runs[0].bold for p in bullets)
