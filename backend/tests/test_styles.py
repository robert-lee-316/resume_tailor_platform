"""Word typography overrides are actually written into DOCX XML."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from app.engine import render_docx


def test_each_text_area_has_independent_style(tmp_path):
    source = tmp_path / 'style-template.docx'
    output = tmp_path / 'out.docx'
    doc = Document()
    doc.add_paragraph('{{ full_name }}')
    doc.add_paragraph('{{ job_title }}')
    doc.add_paragraph('SUMMARY')
    doc.add_paragraph('{{ summary }}')
    doc.add_paragraph('TECHNICAL SKILLS')
    for text in ('{%p for category in skill_categories %}', '{{ category.name }}: {{ category.skills_text }}', '{%p endfor %}'):
        doc.add_paragraph(text)
    doc.add_paragraph('WORK EXPERIENCE')
    doc.add_paragraph('{%p for exp in experiences %}')
    doc.add_paragraph('{{ exp.job_title }}')
    doc.add_paragraph('{%p for bullet in exp.bullets %}')
    doc.add_paragraph('{{ bullet.text }}')
    doc.add_paragraph('{%p endfor %}')
    doc.add_paragraph('{%p endfor %}')
    doc.save(source)
    context = {
        'full_name': 'Alex Morgan', 'job_title': 'Engineer', 'summary': 'Built **secure APIs**.',
        'skill_categories': [{'name': 'Languages', 'skills_text': 'Python, Java'}],
        'experiences': [{'job_title': 'Developer', 'bullets': [{'text': 'Wrote **Python**.'}]}],
        '_styles': {
            'name': {'font':'Georgia','size':22,'color':'#3377AA','bold':True,'align':'center'},
            'role': {'font':'Arial','size':12,'color':'#000000','bold':False,'align':'right'},
            'summary': {'font':'Calibri','size':11,'color':'#333333','bold':False,'align':'left'},
            'skills': {'font':'Verdana','size':10,'color':'#444444','bold':False,'align':'justify'},
            'experience': {'font':'Arial','size':12,'color':'#222222','bold':False,'align':'left'},
            'body': {'font':'Aptos','size':9,'color':'#111111','bold':False,'align':'left'},
        }
    }
    render_docx(source, output, context)
    paras = Document(output).paragraphs
    assert paras[0].text == 'Alex Morgan'
    assert paras[0].runs[0].font.name == 'Georgia'
    assert paras[0].runs[0].font.size.pt == 22
    assert paras[0].runs[0].bold is True
    assert paras[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paras[1].runs[0].font.name == 'Arial'
    assert paras[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    skills = next(p for p in paras if p.text.startswith('Languages:'))
    assert skills.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert skills.runs[0].font.name == 'Verdana'
    summ = next(p for p in paras if p.text.startswith('Built '))
    assert any(r.text == 'secure APIs' and r.bold for r in summ.runs)
    bullet = next(p for p in paras if p.text.startswith('Wrote'))
    assert bullet.runs[0].font.name == 'Aptos'
    assert any(r.text == 'Python' and r.bold for r in bullet.runs)


def test_invalid_styles_do_not_break_template(tmp_path):
    template=tmp_path/'base.docx'; out=tmp_path/'out.docx'
    d=Document(); d.add_paragraph('{{ full_name }}'); d.save(template)
    render_docx(template,out,{'full_name':'Alex','_styles':{'name':{'font':'RandomFont','size':999,'color':'red','align':'bad'}}})
    assert Document(out).paragraphs[0].text=='Alex'


def test_word_run_properties_precede_text_for_pdf_conversion(tmp_path):
    """LibreOffice ignores font overrides when w:rPr appears after w:t."""
    src = tmp_path / 'template.docx'
    out = tmp_path / 'rendered.docx'
    doc = Document()
    doc.add_paragraph('{{ full_name }}')
    doc.add_paragraph('TECHNICAL SKILLS')
    doc.add_paragraph('{%p for category in skill_categories %}')
    doc.add_paragraph('{{ category.name }}: {{ category.skills_text }}')
    doc.add_paragraph('{%p endfor %}')
    doc.save(src)
    render_docx(src, out, {
        'full_name': 'Alex Morgan',
        'skill_categories': [{'name': 'Languages', 'skills_text': 'Python, TypeScript'}],
        '_styles': {
            'name': {'font': 'Arial', 'size': 24, 'color': '#2266AA', 'align': 'center'},
            'skills': {'font': 'Georgia', 'size': 14, 'color': '#112233', 'align': 'left'},
        },
    })
    rendered = Document(out)
    for paragraph in rendered.paragraphs:
        for run in paragraph._p.r_lst:
            if run.find(qn('w:rPr')) is not None:
                assert run[0].tag == qn('w:rPr'), f'Malformed run: {paragraph.text}'
    assert rendered.paragraphs[0].runs[0].font.name == 'Arial'
    assert rendered.paragraphs[0].runs[0].font.size.pt == 24
    skills = next(p for p in rendered.paragraphs if p.text.startswith('Languages:'))
    assert skills.runs[0].font.name == 'Georgia'
    assert skills.runs[0].font.size.pt == 14
