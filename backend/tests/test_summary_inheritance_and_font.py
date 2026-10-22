from docx import Document
from app.engine import render_docx


def test_summary_text_regular_even_if_placeholder_is_bold(tmp_path):
    source = tmp_path / "bold-template.docx"
    output = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("SUMMARY")
    p = doc.add_paragraph()
    r = p.add_run("{{ summary }}")
    r.bold = True
    doc.save(source)
    render_docx(source, output, {"summary": "I build **React** applications with Python."})
    p = Document(output).paragraphs[1]
    assert p.text == "I build React applications with Python."
    assert any("React" in r.text and r.bold is True for r in p.runs)
    assert all(r.bold is False for r in p.runs if "React" not in r.text and r.text.strip())


def test_plain_summary_stays_regular_even_with_bold_template(tmp_path):
    source = tmp_path / "bold-template.docx"
    output = tmp_path / "out.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{ summary }}").bold = True
    doc.save(source)
    render_docx(source, output, {"summary": "Regular summary content."})
    assert Document(output).paragraphs[0].runs[0].bold is False


def test_legacy_full_summary_bold_override_is_ignored(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("{{ summary }}")
    doc.save(source)
    render_docx(source, output, {"summary": "Make this bold.", "_styles": {
        "summary": {"enabled": True, "changed": ["bold"], "bold": True}
    }})
    assert all(run.bold is False for run in Document(output).paragraphs[0].runs)
