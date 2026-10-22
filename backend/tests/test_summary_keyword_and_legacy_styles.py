from docx import Document
from app.engine import render_docx


def test_saved_bold_paragraph_override_does_not_bold_summary(tmp_path):
    template = tmp_path / "source.docx"
    output = tmp_path / "result.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{ summary }}").bold = True
    doc.save(template)
    render_docx(template, output, {
        "summary": "Build **React** apps with Python.",
        "_styles": {"summary": {"enabled": True, "changed": ["font", "bold"], "font": "Arial", "bold": True}}
    })
    runs = Document(output).paragraphs[0].runs
    assert "".join(r.text for r in runs) == "Build React apps with Python."
    assert all(r.bold is False for r in runs if "React" not in r.text and r.text.strip())
    assert any("React" in r.text and r.bold is True for r in runs)
    assert all(r.font.name == "Arial" for r in runs)


def test_plain_summary_remains_unbolded_with_saved_legacy_override(tmp_path):
    template, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    doc.add_paragraph("{{ summary }}")
    doc.save(template)
    render_docx(template, output, {"summary": "This is plain summary text.", "_styles": {
        "summary": {"enabled": True, "changed": ["bold"], "bold": True}
    }})
    runs = Document(output).paragraphs[0].runs
    assert runs and all(r.bold is False for r in runs)
