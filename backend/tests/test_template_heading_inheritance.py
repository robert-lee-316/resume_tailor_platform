from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from app.engine import render_docx


def create_template(path):
    doc = Document()
    doc.add_paragraph('{{ full_name }}')
    for heading in ('SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(heading)
        r.font.bold = True
        r.font.size = Pt(15)  # Heading formatting supplied by the uploaded Word template
    doc.save(path)


def render(tmp_path, styles):
    template = tmp_path / 'source.docx'
    output = tmp_path / 'output.docx'
    create_template(template)
    render_docx(template, output, {'full_name': 'Test Person', '_styles': styles})
    return {p.text: p for p in Document(output).paragraphs}


def test_frontend_defaults_do_not_override_user_template(tmp_path):
    result = render(tmp_path, {key: {
        'enabled': False, 'font': 'Times New Roman', 'size': 11,
        'bold': True, 'align': 'left', 'color': '#161616'
    } for key in ('skillsHeading', 'experienceHeading', 'educationHeading')})
    for key in ('SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
        assert result[key].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert result[key].runs[0].font.bold is True
        assert result[key].runs[0].font.size.pt == 15  # Original Word formatting is preserved


def test_user_size_only_override_preserves_center_and_bold(tmp_path):
    result = render(tmp_path, {'skillsHeading': {
        'enabled': True, 'changed': ['size'], 'font': 'Arial',
        'size': 16, 'bold': False, 'align': 'left'
    }})
    paragraph = result['SKILLS']
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.runs[0].font.bold is True
    assert paragraph.runs[0].font.size.pt == 16


def test_full_explicit_override_takes_precedence(tmp_path):
    result = render(tmp_path, {'experienceHeading': {
        'enabled': True, 'changed': ['bold', 'align', 'size'],
        'size': 13, 'bold': False, 'align': 'left'
    }})
    paragraph = result['WORK EXPERIENCE']
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraph.runs[0].font.bold is False
    assert paragraph.runs[0].font.size.pt == 13


def test_enabled_style_with_empty_changed_list_does_not_override_template(tmp_path):
    # Regression: migration could leave enabled=true but no changed fields.
    result = render(tmp_path, {key: {
        'enabled': True, 'changed': [], 'font': 'Arial', 'size': 10,
        'bold': False, 'align': 'left', 'color': '#161616'
    } for key in ('skillsHeading', 'experienceHeading', 'educationHeading')})
    for key in ('SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
        assert result[key].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert result[key].runs[0].font.bold is True
        assert result[key].runs[0].font.size.pt == 15


def test_explicit_all_heading_preset_overrides_bad_source_template(tmp_path):
    template = tmp_path / 'plain.docx'
    doc = Document()
    for heading in ('SKILLS', 'WORK EXPERIENCE', 'EDUCATION'):
        paragraph = doc.add_paragraph(heading)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.runs[0]
        run.bold = False
        run.font.size = Pt(10)
    doc.save(template)
    output = tmp_path / 'out.docx'
    payload = {key: {'enabled': True, 'changed': ['size', 'bold', 'align'],
        'size': 15, 'bold': True, 'align': 'center'}
        for key in ('skillsHeading', 'experienceHeading', 'educationHeading')}
    render_docx(template, output, {'_styles': payload})
    for paragraph in Document(output).paragraphs:
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert paragraph.runs[0].bold is True
        assert paragraph.runs[0].font.size.pt == 15
